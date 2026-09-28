"""週1回の金庫確認表(印刷用)のExcelを生成するスクリプト。

python3 tools/build_safe_check_sheet.py で docs/safe-check-sheet.xlsx を出力する。

金庫の中身は3つに分けて数える(gas-cash/Config.gs と同じ区分):
- 釣銭ポーチ: 1万円×5個(固定)
- 金庫金: 購入の支払い用。基準額10万円
- 売上封筒: 現金売上の茶封筒。小銭と補充分は金庫金へ振替し、残りを入金する
「帳簿の金額」は、GASが週次で送る「今あるはずの金額」メールから書き写す。
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

OUT = "docs/safe-check-sheet.xlsx"
FONT = "Arial"
POUCH_COUNT = 5
POUCH_AMOUNT = 10000
FUND_BASE = 100000
DENOMS = [10000, 5000, 1000, 500, 100, 50, 10, 5, 1]  # 2,000円札は使わない
ENVELOPE_ROWS = 8

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOOK_FILL = PatternFill("solid", fgColor="DDEBF7")  # メールから書き写す欄
YEN = '#,##0;[Red]-#,##0;""'  # 0は印刷時に空欄に見せる
YEN_DIFF = '#,##0;[Red]-#,##0;0'  # 差額は0も表示する

# 記入例(tests/ledger.test.js の数字と同じ: 金庫金90,160円、封筒2件74,400円)
EXAMPLE = {
    "date": "2026/10/05(月)", "counter": "山田", "witness": "佐藤",
    "pouches": [10000] * POUCH_COUNT,
    "fund_book": 90160,
    "counts": {10000: 7, 5000: 2, 1000: 8, 500: 2, 100: 10, 50: 2, 10: 5, 5: 2, 1: 0},
    "envelopes": [("1", "9/15", "A社(ケータリング)", 30000, 30000),
                  ("2", "10/4", "B社(オードブル)", 44400, 44400)],
    "total_book": 214560,
}


def f(size=10, bold=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def put(ws, ref, value=None, *, size=10, bold=False, fill=None, align="left",
        fmt=None, border=True, color="000000", wrap=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = f(size, bold, color)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    return c


def merge(ws, rng, value=None, **kw):
    if kw.get("border", True):
        for row in ws[rng]:
            for c in row:
                c.border = BOX
                if kw.get("fill"):
                    c.fill = kw["fill"]
    ws.merge_cells(rng)
    return put(ws, rng.split(":")[0], value, **kw)


def section(ws, row, text):
    merge(ws, f"A{row}:H{row}", text, bold=True, border=False, size=11)
    ws.row_dimensions[row].height = 20


def build_sheet(ws, example=False):
    ex = EXAMPLE if example else None
    for col, w in zip("ABCDEFGH", [11, 10, 12, 12, 11, 12, 12, 10]):
        ws.column_dimensions[col].width = w

    merge(ws, "A1:H1", "金庫確認表(週1回)" + ("  ※記入例" if example else ""),
          size=15, bold=True, align="center", border=False)
    ws.row_dimensions[1].height = 28

    put(ws, "A3", "実施日", bold=True, fill=HEAD_FILL, align="center")
    merge(ws, "B3:C3", ex["date"] if ex else None, fill=INPUT_FILL, align="center")
    put(ws, "D3", "数えた人", bold=True, fill=HEAD_FILL, align="center")
    put(ws, "E3", ex["counter"] if ex else None, fill=INPUT_FILL, align="center")
    put(ws, "F3", "立会い", bold=True, fill=HEAD_FILL, align="center")
    merge(ws, "G3:H3", ex["witness"] if ex else None, fill=INPUT_FILL, align="center")
    ws.row_dimensions[3].height = 26

    merge(ws, "A4:H4", "青い欄は、月曜朝の「【金庫確認】あるはずの金額」メール(または「現金残高」シート)から書き写してから数える。"
          "黄色の欄は数えた結果を書く。", size=8, border=False, color="595959", wrap=True)

    # --- 1. 釣銭ポーチ ---
    section(ws, 6, f"1. 釣銭ポーチ({POUCH_AMOUNT:,}円 × {POUCH_COUNT}個)")
    for col, h in zip("ABCD", ["ポーチ", "実際の金額", "差額", "確認"]):
        put(ws, f"{col}7", h, bold=True, fill=HEAD_FILL, align="center", size=9)
    p0 = 8
    for i in range(POUCH_COUNT):
        row = p0 + i
        put(ws, f"A{row}", f"No.{i + 1}", align="center")
        put(ws, f"B{row}", ex["pouches"][i] if ex else None, fill=INPUT_FILL,
            align="right", fmt=YEN)
        put(ws, f"C{row}", f'=IF(B{row}="","",B{row}-{POUCH_AMOUNT})', align="right",
            fmt=YEN_DIFF)
        put(ws, f"D{row}", "✓" if ex else "□", align="center")
        ws.row_dimensions[row].height = 19
    p1 = p0 + POUCH_COUNT - 1
    ptot = p1 + 1
    put(ws, f"A{ptot}", "合計", bold=True, fill=HEAD_FILL, align="center")
    put(ws, f"B{ptot}", f"=SUM(B{p0}:B{p1})", bold=True, align="right", fmt=YEN)
    put(ws, f"C{ptot}", f'=IF(COUNT(B{p0}:B{p1})=0,"",B{ptot}-{POUCH_COUNT * POUCH_AMOUNT})',
        bold=True, align="right", fmt=YEN_DIFF)
    merge(ws, f"E7:H{ptot}",
          "・ポーチは1個ずつ数え、1万円ちょうどか確認する\n"
          "・現場に持ち出し中のポーチは「持出中」と書き、戻ったら数える\n"
          "・1万円を超えた分は売上封筒へ、足りない分は原因を確認\n"
          "・ポーチのお金を支払いに使わない",
          size=8, wrap=True, border=False, color="595959")

    # --- 2. 金庫金 ---
    r = ptot + 2
    section(ws, r, f"2. 金庫金(支払い用・基準額 {FUND_BASE:,}円)")
    r += 1
    for col, h in zip("ABCDEF", ["金種", "枚数", "金額(円)", "金種", "枚数", "金額(円)"]):
        put(ws, f"{col}{r}", h, bold=True, fill=HEAD_FILL, align="center", size=9)
    left, right = DENOMS[:5], DENOMS[5:]
    d0 = r + 1
    for i in range(5):
        row = d0 + i
        pairs = [(left[i], "A", "B", "C")]
        if i < len(right):
            pairs.append((right[i], "D", "E", "F"))
        for denom, dc, nc, ac in pairs:
            put(ws, f"{dc}{row}", f"{denom:,}円" + ("札" if denom >= 1000 else "玉"),
                align="center", size=9)
            put(ws, f"{nc}{row}", ex["counts"][denom] if ex else None, fill=INPUT_FILL,
                align="right", fmt='#,##0;-#,##0;""')
            put(ws, f"{ac}{row}", f"={denom}*{nc}{row}", align="right", fmt=YEN)
        ws.row_dimensions[row].height = 19
    d1 = d0 + 4
    counts = f"B{d0}:B{d1},E{d0}:E{d0 + len(right) - 1}"
    fr = d1 + 1
    put(ws, f"G{d0}", "実際の合計", bold=True, fill=HEAD_FILL, align="center", size=9)
    merge(ws, f"H{d0}:H{d0}", f'=IF(COUNT({counts})=0,"",SUM(C{d0}:C{d1})+SUM(F{d0}:F{d0 + len(right) - 1}))',
          bold=True, align="right", fmt=YEN)
    put(ws, f"G{d0 + 1}", "帳簿の金額", bold=True, fill=HEAD_FILL, align="center", size=9)
    put(ws, f"H{d0 + 1}", ex["fund_book"] if ex else None, fill=BOOK_FILL, align="right",
        fmt=YEN, bold=True)
    put(ws, f"G{d0 + 2}", "差額", bold=True, fill=HEAD_FILL, align="center", size=9)
    put(ws, f"H{d0 + 2}", f'=IF(OR(H{d0}="",H{d0 + 1}=""),"",H{d0}-H{d0 + 1})',
        bold=True, align="right", fmt=YEN_DIFF)
    put(ws, f"G{d0 + 3}", "補充が必要", bold=True, fill=HEAD_FILL, align="center", size=9)
    put(ws, f"H{d0 + 3}", f'=IF(H{d0}="","",MAX(0,{FUND_BASE}-H{d0}))', align="right", fmt=YEN)
    fund_actual, fund_book = f"H{d0}", f"H{d0 + 1}"

    # --- 3. 売上封筒 ---
    r = fr + 1
    section(ws, r, "3. 売上封筒(茶封筒を1件ずつ。帳簿の金額はメールの一覧から書き写す)")
    r += 1
    for col, h in zip("ABCDEFGH", ["No.", "受取日", "相手先", "", "帳簿の金額", "実際の金額",
                                   "差額", "開封確認"]):
        if col == "C":
            merge(ws, f"C{r}:D{r}", h, bold=True, fill=HEAD_FILL, align="center", size=9)
        elif col != "D":
            put(ws, f"{col}{r}", h, bold=True, fill=HEAD_FILL, align="center", size=9)
    e0 = r + 1
    for i in range(ENVELOPE_ROWS):
        row = e0 + i
        s = ex["envelopes"][i] if ex and i < len(ex["envelopes"]) else None
        put(ws, f"A{row}", i + 1, align="center", size=9)
        put(ws, f"B{row}", s[1] if s else None, fill=BOOK_FILL, align="center", size=9)
        merge(ws, f"C{row}:D{row}", s[2] if s else None, fill=BOOK_FILL, size=9)
        put(ws, f"E{row}", s[3] if s else None, fill=BOOK_FILL, align="right", fmt=YEN)
        put(ws, f"F{row}", s[4] if s else None, fill=INPUT_FILL, align="right", fmt=YEN)
        put(ws, f"G{row}", f'=IF(OR(E{row}="",F{row}=""),"",F{row}-E{row})', align="right",
            fmt=YEN_DIFF)
        put(ws, f"H{row}", "✓" if s else "□", align="center")
        ws.row_dimensions[row].height = 19
    e1 = e0 + ENVELOPE_ROWS - 1
    etot = e1 + 1
    merge(ws, f"A{etot}:D{etot}", "合計", bold=True, fill=HEAD_FILL, align="center")
    put(ws, f"E{etot}", f"=SUM(E{e0}:E{e1})", bold=True, align="right", fmt=YEN)
    put(ws, f"F{etot}", f"=SUM(F{e0}:F{e1})", bold=True, align="right", fmt=YEN)
    put(ws, f"G{etot}", f'=IF(COUNT(F{e0}:F{e1})=0,"",F{etot}-E{etot})', bold=True,
        align="right", fmt=YEN_DIFF)
    put(ws, f"H{etot}", None)

    # --- 4. 全体 ---
    r = etot + 2
    section(ws, r, "4. 金庫全体(1+2+3)")
    r += 1
    rows = [
        ("実際に数えた合計", f'=IF(OR(COUNT(B{p0}:B{p1})=0,{fund_actual}=""),"",B{ptot}+{fund_actual}+F{etot})', None),
        ("あるはずの合計(メールの1行目)", ex["total_book"] if ex else None, BOOK_FILL),
        ("差額", f'=IF(OR(D{r}="",D{r + 1}=""),"",D{r}-D{r + 1})', None),
    ]
    for k, (label, val, fill) in enumerate(rows):
        row = r + k
        merge(ws, f"A{row}:C{row}", label, bold=True, fill=HEAD_FILL, size=9)
        put(ws, f"D{row}", val, bold=True, align="right", fill=fill,
            fmt=YEN_DIFF if k == 2 else YEN)
        ws.row_dimensions[row].height = 20
    merge(ws, f"E{r}:H{r + 2}",
          "差額が1円でもあれば下に原因を書き、当日中に責任者へ報告。\n"
          "原因が分かったら、フォームの登録漏れを追加登録する。",
          size=8, wrap=True, border=False, color="595959")

    # --- 5. 確認項目 ---
    r = r + len(rows) + 1
    section(ws, r, "5. 確認項目(はい/いいえ に○)")
    checks = [
        "受け取りから14日を超えた売上封筒はない(あれば次に銀行へ行くとき必ず入金)",
        "すべての封筒に 受取日・相手先・金額・担当者 が書いてある",
        "今週の支払いのレシートが全部そろっている",
        "売上封筒から金庫金へ移したお金は、すべて「振替」で登録してある",
        "数えた後、2人で施錠を確認した",
    ]
    for i, text in enumerate(checks):
        row = r + 1 + i
        merge(ws, f"A{row}:F{row}", f"□ {text}", size=9)
        merge(ws, f"G{row}:H{row}", "はい ・ いいえ", align="center", size=9)
        ws.row_dimensions[row].height = 18
    r = r + len(checks) + 2
    put(ws, f"A{r}", "差額の原因\n・対応", bold=True, fill=HEAD_FILL, align="center", size=9,
        wrap=True)
    merge(ws, f"B{r}:F{r}", None, fill=INPUT_FILL, size=9, wrap=True)
    put(ws, f"G{r}", "責任者印", bold=True, fill=HEAD_FILL, align="center", size=9)
    put(ws, f"H{r}", None)
    ws.row_dimensions[r].height = 40
    last = r

    ws.print_area = f"A1:H{last}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.5, bottom=0.5)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "記入後は責任者が押印し、ファイルに綴じて1年間保管"
    ws.oddFooter.center.size = 8
    ws.sheet_view.showGridLines = False


def build_guide(ws):
    ws.column_dimensions["A"].width = 110
    lines = [
        ("金庫確認表の使い方", True),
        ("", False),
        ("金庫の中身は3つに分ける", True),
        (f"・釣銭ポーチ: {POUCH_AMOUNT:,}円 × {POUCH_COUNT}個。現場に持っていくお釣り。支払いには使わない。", False),
        (f"・金庫金: 購入の現金払いに使うお金。基準額 {FUND_BASE:,}円。減った分は銀行から引き出して補充する。", False),
        ("・売上封筒: ケータリング・オードブルの現金売上。茶封筒に入れて保管する。入金の日に小銭と補充分を金庫金へ振替し、残りのお札を入金する。", False),
        ("", False),
        ("毎回(お金を出し入れしたとき)", True),
        ("・紙には書かず、スマホのフォーム(大阪GB 現金出納帳)から登録する。処理区分は5つから選ぶ。", False),
        ("・売上封筒には、受取日・相手先・金額・担当者を書く。支払いには使わない。", False),
        ("", False),
        ("週1回(月曜)", True),
        ("1. 月曜朝に届くメール「【金庫確認】あるはずの金額」の数字を、この表の青い欄に書き写す。", False),
        ("2. 2人で、ポーチ・金庫金・売上封筒の順に数え、黄色の欄に書く。", False),
        ("3. 差額が1円でもあれば原因を書き、当日中に責任者へ報告。登録漏れならフォームで追加登録する。", False),
        ("4. 責任者が押印し、ファイルに綴じて1年間保管する。", False),
        ("", False),
        ("ルール", True),
        ("・数える人は毎週交代する(同じ人に固定しない)。", False),
        ("・売上封筒のお金で直接支払わない。金庫金へ移すときは「振替(売上封筒から金庫金へ)」で登録する。", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = f(14 if i == 1 else 11, bold)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.sheet_view.showGridLines = False


wb = Workbook()
guide = wb.active
guide.title = "使い方"
build_guide(guide)
build_sheet(wb.create_sheet("確認表(印刷用)"))
build_sheet(wb.create_sheet("記入例"), example=True)
wb.active = 1
wb.save(OUT)
print("saved", OUT)
